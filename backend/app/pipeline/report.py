"""PDF and Excel exports of a finished check (TZ §9)."""

import io
from pathlib import Path

from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.models import Match, Project, Recommendation, Run

STATUS = {
    "ru": {"match": "Соответствует", "partial": "Частично", "unconfirmed": "Не подтверждено",
           "mismatch": "Не соответствует", "review": "Требует проверки"},
    "uz": {"match": "Mos", "partial": "Qisman", "unconfirmed": "Tasdiqlanmagan",
           "mismatch": "Mos emas", "review": "Tekshirish kerak"},
    "en": {"match": "Compliant", "partial": "Partial", "unconfirmed": "Not confirmed",
           "mismatch": "Non-compliant", "review": "Needs review"},
}  # fmt: skip
COLORS = {"match": "16A34A", "partial": "EAB308", "unconfirmed": "F97316",
          "mismatch": "DC2626", "review": "6B7280"}  # fmt: skip
TYPES = {
    "ru": {"mandatory": "Обязательное", "rated": "Оценочное", "desired": "Желательное"},
    "uz": {"mandatory": "Majburiy", "rated": "Baholanadigan", "desired": "Istalgan"},
    "en": {"mandatory": "Mandatory", "rated": "Rated", "desired": "Desired"},
}
TEXT = {
    "ru": {"title": "Отчёт о соответствии заявки", "total": "Соответствие", "product": "Товар",
           "participant": "Участник", "mandatory": "Обязательные выполнены", "verdict": "Вердикт",
           "forecast": "После устранения", "summary": "Аналитический вывод",
           "plan": "План действий", "table": "Сравнение по требованиям", "req": "Требование",
           "bid": "Заявка", "status": "Статус", "reason": "Обоснование", "action": "Действие",
           "impact": "Влияние", "warnings": "Предупреждения", "deadline": "Подача до",
           "requirements": "Требования", "facts": "Факты заявки", "source": "Источник", "type": "Тип", "quote": "Цитата"},
    "uz": {"title": "Ariza muvofiqligi hisoboti", "total": "Moslik", "product": "Tovar",
           "participant": "Ishtirokchi", "mandatory": "Bajarilgan majburiylar",
           "verdict": "Verdikt", "forecast": "Tuzatilgandan keyin",
           "summary": "Analitik xulosa", "plan": "Harakatlar rejasi",
           "table": "Talablar bo'yicha solishtirish", "req": "Talab", "bid": "Ariza",
           "status": "Status", "reason": "Asos", "action": "Harakat", "impact": "Ta'siri",
           "warnings": "Ogohlantirishlar", "deadline": "Topshirish muddati",
           "requirements": "Talablar", "facts": "Ariza faktlari", "source": "Manba", "type": "Turi", "quote": "Iqtibos"},
    "en": {"title": "Bid compliance report", "total": "Compliance", "product": "Product",
           "participant": "Participant", "mandatory": "Mandatory met", "verdict": "Verdict",
           "forecast": "After fixes", "summary": "Analytical conclusion",
           "plan": "Action plan", "table": "Requirement-by-requirement comparison",
           "req": "Requirement", "bid": "Bid", "status": "Status", "reason": "Reasoning",
           "action": "Action", "impact": "Impact", "warnings": "Warnings",
           "deadline": "Submission deadline", "requirements": "Requirements",
           "facts": "Bid facts", "source": "Source", "type": "Type", "quote": "Quote"},
}  # fmt: skip
SEVERITY_ORDER = {"high": 0, "medium": 1, "low": 2}

_templates = Environment(
    loader=FileSystemLoader(Path(__file__).resolve().parent.parent / "templates"),
    autoescape=select_autoescape(["html"]),
)


def sorted_matches(run: Run) -> list[Match]:
    return sorted(
        run.matches, key=lambda m: (m.requirement.scope != "product", m.requirement.sort_index)
    )


def sorted_recommendations(run: Run) -> list[Recommendation]:
    return sorted(run.recommendations, key=lambda r: (SEVERITY_ORDER[r.severity], -r.impact))


def condition(match: Match) -> str:
    req = match.requirement
    sign = {"gte": "≥", "lte": "≤", "eq": "=", "range_within": "∈", "range_covers": "⊇"}
    parts = [req.parameter]
    if req.value:
        parts += [sign.get(req.operator, ":"), req.value, req.unit or ""]
    return " ".join(p for p in parts if p).strip()


def bid_value(match: Match) -> str:
    return "; ".join(e["value"] for e in match.evidence if e.get("value")) or "—"


def build_pdf(project: Project, run: Run) -> bytes:
    from weasyprint import HTML  # heavy import, only needed for this export

    lang = project.report_language
    html = _templates.get_template("report.html").render(
        project=project,
        run=run,
        score=run.score,
        matches=sorted_matches(run),
        recommendations=sorted_recommendations(run),
        t=TEXT[lang],
        status_label=STATUS[lang],
        type_label=TYPES[lang],
        colors=COLORS,
        condition=condition,
        bid_value=bid_value,
    )
    return HTML(string=html).write_pdf()  # type: ignore[no-any-return]


def build_xlsx(project: Project, run: Run) -> bytes:
    lang = project.report_language
    t, labels = TEXT[lang], STATUS[lang]
    score = run.score
    book = Workbook()
    bold = Font(bold=True)

    summary = book.active
    assert summary is not None
    summary.title = t["summary"][:31]
    rows = [
        (t["title"], project.title),
        (t["deadline"], project.submission_deadline.isoformat() if project.submission_deadline else ""),
        (t["verdict"], score.get("verdict_label")),
        (t["total"], score.get("total")),
        (t["product"], score.get("product")),
        (t["participant"], score.get("participant")),
        (t["mandatory"], f"{score.get('mandatory_done')} / {score.get('mandatory_total')}"),
        (t["forecast"], score.get("forecast")),
        (t["warnings"], "; ".join(score.get("warnings", []))),
        (t["summary"], run.summary or ""),
    ]  # fmt: skip
    for row in rows:
        summary.append(row)
    for cell in summary["A"]:
        cell.font = bold
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 100
    summary["B10"].alignment = Alignment(wrap_text=True, vertical="top")

    compare = book.create_sheet(t["table"][:31])
    compare.append(["ID", t["req"], "", t["bid"], t["status"], t["reason"], t["source"]])
    for match in sorted_matches(run):
        sources = "; ".join(
            f"{e['filename']} p.{e.get('page') or '?'}" for e in match.evidence if e.get("filename")
        )
        compare.append([
            match.requirement.code, condition(match), TYPES[lang][match.requirement.req_type],
            bid_value(match), labels[match.status], match.reasoning, sources,
        ])  # fmt: skip
        fill = PatternFill("solid", fgColor=COLORS[match.status])
        compare.cell(row=compare.max_row, column=5).fill = fill
        compare.cell(row=compare.max_row, column=5).font = Font(color="FFFFFF", bold=True)
    for col, width in zip("ABCDEFG", (9, 50, 15, 30, 18, 70, 40), strict=True):
        compare.column_dimensions[col].width = width
    for cell in compare[1]:
        cell.font = bold

    plan = book.create_sheet(t["plan"][:31])
    plan.append(["ID", t["req"], t["reason"], t["action"], t["type"], t["impact"], "✓"])
    for rec in sorted_recommendations(run):
        plan.append([
            rec.requirement.code, rec.requirement.parameter, rec.problem, rec.action,
            rec.action_type, f"+{rec.impact}%", "✓" if rec.done else "",
        ])  # fmt: skip
    for col, width in zip("ABCDEFG", (9, 35, 60, 60, 22, 10, 5), strict=True):
        plan.column_dimensions[col].width = width
    for cell in plan[1]:
        cell.font = bold

    reqs = book.create_sheet(t["requirements"][:31])
    reqs.append(["ID", t["req"], t["type"], t["source"], t["quote"]])
    for match in sorted_matches(run):
        src = match.requirement.sources[0] if match.requirement.sources else {}
        reqs.append([
            match.requirement.code, condition(match), TYPES[lang][match.requirement.req_type],
            f"{src.get('filename', '')} p.{src.get('page') or '?'}", src.get("quote") or "",
        ])  # fmt: skip
    for col, width in zip("ABCDE", (9, 50, 15, 40, 80), strict=True):
        reqs.column_dimensions[col].width = width
    for cell in reqs[1]:
        cell.font = bold

    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()
