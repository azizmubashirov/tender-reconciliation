"""A services tender checked against a bid merged into one PDF (the case of project «Физкультура»).

The model is scripted; the test covers what code decides around it: document sections inside one
file, self-declared terms, money units, document numbers, "not applicable" items, evaluation
points against the passing score, and a summary written with the final numbers.
"""

import io
import re
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient
from pydantic import BaseModel

from app.models import Document
from app.pipeline.check import _apply_doc_results, _doc_sections, _section_at
from app.schemas.llm import (
    BidDocResult,
    BidFactOut,
    BidSectionOut,
    CustomerDocResult,
    ExtractedRequirement,
    JudgeResult,
    LotPositionOut,
    RecommendationOut,
    ReportResult,
    SummaryResult,
    VerdictOut,
)
from tests.conftest import ClientFactory, make_pdf

TENDER = [
    "Lot 26521012040746. Services: write-off and disposal of fixed assets. Start price "
    "120 000 000 so'm with VAT. The price is entered in the electronic system.",
    "Service term shall not exceed 7 days. The bidder must have at least one certified technical "
    "specialist (engineer-expert). Statement of the Tax Committee on the absence of tax debt.",
    "Recommendation letters: 3 letters - 2 points, 1 to 3 letters - 1 point, none - 0. "
    "Experience 2025: contracts over 500 mln so'm - 3 points, 300-500 mln - 2 points, less than "
    "300 mln - 1 point. Contract with a recycling plant: present - 4, absent - 0. "
    "Maximum score 9, minimum passing score 7.",
]
BID = [
    "Form 6. Terms of service: not more than 7 (seven) days. The price is entered in the "
    "electronic system.",
    "MY.SOLIQ.UZ Tax Committee. Statement No 982172. No tax debt. Valid until 2026-11-01.",
    "Recommendation letter BUNYOD. Recommendation letter OBOD. Recommendation letter MOVE ON.",
    "Invoice No 4 dated 03.07.2025 to contract No 8/02. Total: 150 112 592.00 sum. "
    "Engineer-expert, certificate No 462945631.",
]


def req(code_hint: str, **kw: Any) -> ExtractedRequirement:
    base: dict[str, Any] = {
        "scope": "participant", "category": "qualification", "lot_position": None,
        "parameter_key": code_hint, "value": None, "unit": None, "req_type": "mandatory",
        "rated_points": None, "proof_documents": [], "consequence": "rejection",
        "equivalent_allowed": False, "restrictive_note": None, "page": None,
    }  # fmt: skip
    return ExtractedRequirement(**(base | kw))


def section(first: int, doc_type: str, title: str, **kw: Any) -> BidSectionOut:
    base: dict[str, Any] = {
        "issuer": None, "doc_number": None, "doc_date": None, "valid_until": None,
        "has_signature": True, "has_seal": True,
    }  # fmt: skip
    return BidSectionOut(
        first_page=first, last_page=first, doc_type=doc_type, title=title, **(base | kw)
    )


def fact(code: str, value: str | None, page: int, quote: str) -> BidFactOut:
    return BidFactOut(requirement_code=code, value=value, unit=None, page=page, quote=quote)


class Model:
    """Scripted Claude: answers by schema and records every prompt it gets."""

    def __init__(self) -> None:
        self.prompts: dict[type[BaseModel], list[str]] = {}

    def __call__(
        self, schema: type[BaseModel], *, system: str, content: list[dict[str, Any]], **_: Any
    ) -> Any:
        text = " ".join(str(b.get("title", "")) + str(b.get("text", "")) for b in content)
        self.prompts.setdefault(schema, []).append(text)
        return getattr(self, schema.__name__)(text)

    def CustomerDocResult(self, text: str) -> CustomerDocResult:  # noqa: N802
        return CustomerDocResult(
            doc_type="tender_documentation", title="Disposal of fixed assets",
            customer_name="University", lot_number="26521012040746",
            submission_deadline="2026-10-05", min_passing_score=7,
            lot_positions=[LotPositionOut(number=1, name="Disposal services", quantity=None,
                                          unit=None)],
            requirements=[
                req("service_term", scope="product", category="commercial",
                    parameter="Service term", operator="lte", value="7", unit="kun",
                    proof_documents=["commercial_proposal"],
                    page=2, quote="Service term shall not exceed 7 days"),
                req("price", category="commercial", parameter="Offer price", operator="lte",
                    value="120 000 000", unit="so‘m (QQS bilan)",
                    proof_documents=["commercial_proposal"], page=1,
                    quote="Start price 120 000 000 so'm with VAT"),
                req("engineer", parameter="Certified engineer-expert", operator="gte", value="1",
                    unit="xodim", proof_documents=["certificate"], page=2,
                    quote="at least one certified technical specialist (engineer-expert)"),
                req("tax", category="documents", parameter="No tax debt", operator="exists",
                    proof_documents=["certificate"], page=2,
                    quote="Statement of the Tax Committee on the absence of tax debt"),
                req("letters", req_type="rated", rated_points=2, parameter="Recommendation letters",
                    operator="gte", value="3", unit="dona", proof_documents=["reference_letter"],
                    consequence="score_reduction", page=3,
                    quote="Recommendation letters: 3 letters - 2 points"),
                req("experience", req_type="rated", rated_points=3, parameter="Contracts 2025",
                    operator="gte", value="500", unit="mln so'm",
                    proof_documents=["experience_contract"], consequence="score_reduction",
                    page=3, quote="contracts over 500 mln so'm - 3 points"),
                req("plant", req_type="rated", rated_points=4, parameter="Contract with a plant",
                    operator="exists", proof_documents=["experience_contract"],
                    consequence="score_reduction", page=3,
                    quote="Contract with a recycling plant: present - 4, absent - 0"),
                req("pass", parameter="Minimal o'tish bali", operator="gte", value="7",
                    unit="ball", page=3, quote="minimum passing score 7"),
            ],
            issues=[],
        )  # fmt: skip

    def BidDocResult(self, text: str) -> BidDocResult:  # noqa: N802
        return BidDocResult(
            doc_type="tech_proposal", participant_name="INFOSISTEMA", participant_tin="203524684",
            sections=[
                section(1, "tech_proposal", "Form 6"),
                section(2, "certificate", "Tax statement", issuer="Tax Committee",
                        valid_until="2026-11-01"),
                section(3, "reference_letter", "Recommendation letters"),
                section(4, "experience_contract", "Invoice No 4"),
            ],
            inconsistencies=["Form 2: registration date 04.02.2011, registry extract: 16.02.2001"],
            facts=[
                fact("T-001", "7 days", 1, "Terms of service: not more than 7 (seven) days"),
                fact("U-001", None, 1, "The price is entered in the electronic system"),
                fact("U-002", "№ 462945631", 4, "Engineer-expert, certificate No 462945631"),
                # The model got the page wrong: the quote is found on page 2.
                fact("U-003", "no tax debt", 1, "Statement No 982172. No tax debt."),
                fact("U-004", "3 шт", 3, "Recommendation letter MOVE ON"),
                fact("U-005", "150 112 592.00 сум", 4, "Total: 150 112 592.00 sum"),
            ],
        )  # fmt: skip

    def JudgeResult(self, text: str) -> JudgeResult:  # noqa: N802
        tax_confirmed = '"bid.pdf" (certificate) p.2' in text
        decided: dict[str, tuple[str, float | None]] = {
            "U-001": ("na", None),
            "U-002": ("review", None),
            "U-003": ("match" if tax_confirmed else "unconfirmed", None),
            "U-004": ("match", 2),
            "U-005": ("mismatch", 1),
            "U-006": ("mismatch", 0),
            "U-007": ("review", None),
        }
        return JudgeResult(verdicts=[
            VerdictOut(requirement_code=code, status=status, reasoning=f"judged {code}",
                       evidence_quote=None, evidence_document=None, awarded_points=points)
            for code in re.findall(r"\[([TU]-\d{3}) \|", text)
            for status, points in [decided[code]]
        ])  # fmt: skip

    def ReportResult(self, text: str) -> ReportResult:  # noqa: N802
        return ReportResult(recommendations=[
            RecommendationOut(requirement_code=code, problem=f"problem {code}",
                              action=f"action {code}",
                              action_type="unfixable" if code == "U-006" else "add_document")
            for code in re.findall(r"\[([TU]-\d{3})\]", text)
        ])  # fmt: skip

    def SummaryResult(self, text: str) -> SummaryResult:  # noqa: N802
        return SummaryResult(summary="Риск отклонения: баллы ниже проходного.")


@pytest.fixture
def model() -> Model:
    return Model()


@pytest.fixture
def client(client_with: ClientFactory, model: Model) -> TestClient:
    return client_with(model)


def upload(client: TestClient, pid: int, kind: str, name: str, pages: list[str]) -> int:
    package = client.post(f"/api/v1/projects/{pid}/packages", json={"kind": kind}).json()
    client.post(
        f"/api/v1/packages/{package['id']}/files",
        files=[("files", (name, make_pdf(pages), "application/pdf"))],
    )
    assert client.post(f"/api/v1/packages/{package['id']}/process").status_code == 200
    return int(package["id"])


def test_merged_bid_pdf(client: TestClient, model: Model) -> None:
    pid = client.post("/api/v1/projects", json={"title": "Физкультура"}).json()["id"]
    upload(client, pid, "customer", "tender.pdf", TENDER)
    project = client.get(f"/api/v1/projects/{pid}").json()
    assert project["min_pass_score"] == 7
    client.post(f"/api/v1/projects/{pid}/checklist/approve")

    bid_id = upload(client, pid, "bid", "bid.pdf", BID)
    run = client.post(f"/api/v1/projects/{pid}/runs", json={"package_id": bid_id}).json()
    run = client.get(f"/api/v1/runs/{run['id']}").json()
    assert run["status"] == "done", run["error"]

    # The single file is a package of four documents, each with its own type.
    doc = client.get(f"/api/v1/projects/{pid}").json()["packages"][1]["documents"][0]
    assert doc["doc_type"] == "bid_package"
    assert [s["doc_type"] for s in doc["info"]["sections"]] == [
        "tech_proposal", "certificate", "reference_letter", "experience_contract"
    ]  # fmt: skip

    judge = model.prompts[JudgeResult][0]
    assert '- "bid.pdf" p.2, certificate, Tax statement, Tax Committee' in judge
    assert 'source: "Recommendation letters: 3 letters - 2 points"' in judge
    assert "code comparison: 3 dona ≥ 3 dona -> match (final)" in judge

    matches = {
        m["requirement"]["code"]: m for m in client.get(f"/api/v1/runs/{run['id']}/matches").json()
    }
    status = {code: (m["status"], m["method"]) for code, m in matches.items()}
    assert status == {
        "T-001": ("match", "rule"),  # own form proves the own term (was "unconfirmed")
        "U-001": ("na", "llm"),  # price lives in the e-procurement system
        "U-002": ("review", "llm"),  # "№ 462945631" is not a head count
        "U-003": ("match", "llm"),  # the tax statement on page 2 confirms it
        "U-004": ("match", "rule"),  # three letters inside the same PDF
        "U-005": ("mismatch", "rule"),  # 150 mln < 500 mln
        "U-006": ("mismatch", "llm"),
        "U-007": ("mismatch", "rule"),  # decided from points, not by the model
    }
    assert matches["T-001"]["reasoning"] == "7 kun ≤ 7 kun"
    assert matches["U-003"]["evidence"][0]["page"] == 2
    assert matches["U-003"]["evidence"][0]["doc_type"] == "certificate"
    assert matches["U-005"]["reasoning"] == "150.113 mln so'm < 500 mln so'm"
    assert [matches[c]["points"] for c in ("U-004", "U-005", "U-006")] == [2, 1, 0]
    assert matches["U-007"]["reasoning"] == "Баллы заявки 3–3 из 9 < проходного балла 7"

    score = run["score"]
    assert score["points"] == {"min": 3, "max": 3, "total": 9, "pass": 7, "passed": False}
    assert score["verdict"] == "rejection_risk"
    assert score["counts"]["na"] == 1
    assert (score["mandatory_done"], score["mandatory_total"]) == (2, 4)
    assert score["points_forecast"] == 5  # the plant contract is unfixable
    assert any("registration date 04.02.2011" in w for w in score["warnings"])

    # The summary is written after the plan, with the final forecast and points.
    summary = model.prompts[SummaryResult][0]
    assert f"recommendations are done: {score['forecast']:g}%" in summary
    assert "Evaluation points: 3–3 of 9; at most 5" in summary
    assert "Passing score 7: NOT reachable" in summary
    assert "Not applicable: [U-001]" in summary
    assert run["summary"] == "Риск отклонения: баллы ниже проходного."

    # Changing the passing score re-decides the threshold item and the verdict.
    client.patch(f"/api/v1/projects/{pid}", json={"min_pass_score": 3})
    after = client.get(f"/api/v1/runs/{run['id']}").json()["score"]
    assert after["points"]["passed"] is True
    assert after["verdict"] == "needs_work"  # U-002 still needs a human
    threshold = next(
        m for m in client.get(f"/api/v1/runs/{run['id']}/matches").json()
        if m["requirement"]["code"] == "U-007"
    )  # fmt: skip
    assert threshold["status"] == "match"

    # A manual "match" gives the full points of the criterion.
    client.patch(
        f"/api/v1/matches/{matches['U-006']['id']}",
        json={"status": "match", "comment": "Договор с заводом приложен"},
    )
    assert client.get(f"/api/v1/runs/{run['id']}").json()["score"]["points"]["min"] == 7

    xlsx = client.get(f"/api/v1/runs/{run['id']}/report.xlsx")
    book = openpyxl.load_workbook(io.BytesIO(xlsx.content))
    labels = [row[4] for row in book.worksheets[1].iter_rows(min_row=2, values_only=True)]
    assert "Не применяется" in labels
    assert book.worksheets[0]["B9"].value == "7 из 9 (проходной балл 3)"
    assert client.get(f"/api/v1/runs/{run['id']}/report.pdf").content.startswith(b"%PDF")


def test_sections_of_a_split_pdf_use_file_pages() -> None:
    """Files over 80 pages are sent in parts: section pages must be counted in the whole file."""
    doc = Document(filename="bid.pdf", storage_key="k", page_count=102)

    def part(doc_type: str, first: int, last: int) -> BidDocResult:
        common: dict[str, Any] = {"participant_name": None, "participant_tin": None, "facts": []}
        sec = section(first, doc_type, doc_type).model_copy(update={"last_page": last})
        return BidDocResult(doc_type=doc_type, **common, sections=[sec], inconsistencies=[])

    _apply_doc_results(
        doc, [(1, part("tech_proposal", 1, 80)), (81, part("reference_letter", 8, 12))]
    )
    sections = _doc_sections(doc)
    assert doc.doc_type == "bid_package"
    assert [(s["first_page"], s["last_page"]) for s in sections] == [(1, 80), (88, 92)]
    found = _section_at(sections, 90)
    assert found is not None and found["doc_type"] == "reference_letter"
