"""End-to-end flow through the API with real file conversion and a scripted model.

Covers the MVP chain: upload -> checklist -> bid facts -> matching -> score -> plan -> exports.
"""

import io
import re
import zipfile
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import fitz
import openpyxl
import pytest
from fastapi.testclient import TestClient
from pydantic import BaseModel
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app import llm
from app.core import db as dbmod
from app.core import storage
from app.core.config import Settings
from app.main import app
from app.schemas.llm import (
    BidDocResult,
    BidFactOut,
    CustomerDocResult,
    CustomerIssueOut,
    ExtractedRequirement,
    JudgeResult,
    LotPositionOut,
    RecommendationOut,
    ReportResult,
    VerdictOut,
)
from app.workers import tasks

TENDER = [
    "Tender No 2026-114. Supply of laptops. Bids are accepted until 15.10.2026.",
    "Technical specification. Position 1: Laptop, 20 pcs. RAM shall be not less than 16 GB. "
    "Warranty period shall be not less than 36 months. The laptop must have an anti-glare display.",
    "Qualification. The bidder must provide a valid ISO 9001 certificate. "
    "Experience: at least 3 similar contracts during the last 2 years.",
]
PROPOSAL = [
    "Technical proposal. Laptop Model X1. RAM 16 GB. Warranty 36 months. Anti-glare matte display."
]
PASSPORT = ["Product passport Model X1. Memory: 16 GB. Warranty: 24 months."]


def make_pdf(pages: list[str]) -> bytes:
    doc = fitz.open()
    for text in pages:
        doc.new_page().insert_textbox(fitz.Rect(50, 50, 550, 800), text, fontsize=11)
    data: bytes = doc.tobytes()
    return data


def req(**kw: Any) -> ExtractedRequirement:
    base: dict[str, Any] = {
        "scope": "product", "category": "technical", "lot_position": 1, "unit": None,
        "value": None, "req_type": "mandatory", "rated_points": None, "proof_documents": [],
        "consequence": "rejection", "equivalent_allowed": False, "restrictive_note": None,
    }  # fmt: skip
    return ExtractedRequirement(**(base | kw))


def fact(code: str, value: str, quote: str) -> BidFactOut:
    return BidFactOut(requirement_code=code, value=value, unit=None, page=1, quote=quote)


def fake_ask(
    schema: type[BaseModel], *, system: str, content: list[dict[str, Any]], **_: Any
) -> Any:
    text = " ".join(str(block.get("title", "")) + str(block.get("text", "")) for block in content)
    if schema is CustomerDocResult:
        return CustomerDocResult(
            doc_type="technical_specification", title="Supply of laptops", customer_name="Ministry",
            lot_number="2026-114", submission_deadline="2026-10-15",
            lot_positions=[LotPositionOut(number=1, name="Laptop", quantity="20", unit="pcs")],
            requirements=[
                req(parameter="RAM", parameter_key="ram", operator="gte", value="16", unit="GB",
                    proof_documents=["product_passport"], page=2,
                    quote="RAM shall be not less than 16 GB"),
                req(parameter="Warranty", parameter_key="warranty", operator="gte", value="36",
                    unit="months", proof_documents=["product_passport"], page=2,
                    quote="Warranty period shall be not less than 36 months"),
                req(parameter="Anti-glare display", parameter_key="anti_glare", operator="text",
                    page=2, quote="The laptop must have an anti-glare display"),
                # Same requirement found twice: merged into one with two sources.
                req(parameter="Оперативная память", parameter_key="RAM", operator="gte",
                    value="16", unit="ГБ", page=2, quote="RAM shall be not less than 16 GB"),
                req(scope="participant", category="documents", lot_position=None,
                    parameter="ISO 9001 certificate", parameter_key="iso_9001", operator="exists",
                    proof_documents=["certificate"], page=1,
                    quote="The bidder must provide a valid ISO 9001 certificate"),
                req(scope="participant", category="qualification", lot_position=None,
                    parameter="Similar contracts", parameter_key="experience", operator="text",
                    req_type="rated", rated_points=10, consequence="score_reduction", page=3,
                    quote="at least 3 similar contracts during the last 2 years"),
            ],
            issues=[CustomerIssueOut(kind="ambiguity", description="Не указан тип матрицы",
                                     quote=None, page=None)],
        )  # fmt: skip
    if schema is BidDocResult:
        common: dict[str, Any] = {
            "title": None, "issuer": None, "doc_number": None, "doc_date": None,
            "valid_until": None, "has_signature": True, "has_seal": True,
            "participant_name": "Tech LLC", "participant_tin": "301234567",
        }  # fmt: skip
        if "proposal.pdf" in text:
            return BidDocResult(doc_type="tech_proposal", **common, facts=[
                fact("T-001", "16 GB", "RAM 16 GB"),
                fact("T-002", "36 months", "Warranty 36 months"),
                fact("T-003", "matte display", "Anti-glare matte display"),
            ])  # fmt: skip
        if "passport.pdf" in text:
            return BidDocResult(doc_type="product_passport", **common, facts=[
                fact("T-001", "16 GB", "Memory: 16 GB"),
                fact("T-002", "24 months", "Warranty: 24 months"),
            ])  # fmt: skip
        return BidDocResult(doc_type="experience_contract", **common, facts=[
            fact("U-002", "2 contracts", "No1,2025-03-01"),
        ])  # fmt: skip
    if schema is JudgeResult:
        decided = {"T-003": "match", "U-001": "mismatch", "U-002": "partial"}
        codes = re.findall(r"\[([TU]-\d{3}) \|", text)
        return JudgeResult(verdicts=[
            VerdictOut(requirement_code=c, status=decided[c], reasoning=f"judged {c}",
                       evidence_quote=None, evidence_document=None)
            for c in codes
        ])  # fmt: skip
    if schema is ReportResult:
        codes = re.findall(r"\[([TU]-\d{3})\]", text)
        return ReportResult(summary="Риск отклонения: гарантия ниже требуемой.", recommendations=[
            RecommendationOut(requirement_code=c, problem=f"problem {c}", action=f"fix {c}",
                              action_type="fix_data")
            for c in codes
        ])  # fmt: skip
    raise AssertionError(schema)


@pytest.fixture
def client(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Iterator[TestClient]:
    engine = create_engine(
        "sqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False}
    )
    sessions = sessionmaker(engine, expire_on_commit=False)
    monkeypatch.setattr(dbmod, "engine", engine)
    monkeypatch.setattr(dbmod, "SessionLocal", sessions)
    monkeypatch.setattr(tasks, "SessionLocal", sessions)
    settings = Settings(local_storage_dir=str(tmp_path), _env_file=None)
    monkeypatch.setattr(storage, "get_settings", lambda: settings)
    monkeypatch.setattr(llm, "ask", fake_ask)
    # Run Celery tasks inline.
    for task in (tasks.process_package, tasks.extract_requirements, tasks.check_run):
        monkeypatch.setattr(task, "delay", task)
    with TestClient(app) as test_client:
        yield test_client


def test_full_check(client: TestClient) -> None:
    project = client.post("/api/v1/projects", json={"title": "Laptops"}).json()
    pid = project["id"]

    customer = client.post(f"/api/v1/projects/{pid}/packages", json={"kind": "customer"}).json()
    client.post(
        f"/api/v1/packages/{customer['id']}/files",
        files=[("files", ("tender.pdf", make_pdf(TENDER), "application/pdf"))],
    )
    assert client.post(f"/api/v1/packages/{customer['id']}/process").status_code == 200

    project = client.get(f"/api/v1/projects/{pid}").json()
    assert project["status"] == "extracted"
    assert project["submission_deadline"] == "2026-10-15"
    assert project["lot_positions"][0]["name"] == "Laptop"

    reqs = client.get(f"/api/v1/projects/{pid}/requirements").json()
    assert [r["code"] for r in reqs] == ["T-001", "T-002", "T-003", "U-001", "U-002"]
    ram = reqs[0]
    assert len(ram["sources"]) == 2  # merged duplicate
    assert ram["sources"][0]["verified"] is True
    assert reqs[3]["sources"][0]["page"] == 3  # model said page 1, quote is on page 3
    assert client.get(f"/api/v1/projects/{pid}/issues").json()[0]["kind"] == "ambiguity"

    assert client.post(f"/api/v1/projects/{pid}/checklist/approve").json()["checklist_approved"]

    bid = client.post(f"/api/v1/projects/{pid}/packages", json={"kind": "bid"}).json()
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as zf:
        zf.writestr("experience/contracts.csv", "contract,date\nNo1,2025-03-01\nNo2,2025-09-10\n")
    client.post(
        f"/api/v1/packages/{bid['id']}/files",
        files=[
            ("files", ("proposal.pdf", make_pdf(PROPOSAL), "application/pdf")),
            ("files", ("passport.pdf", make_pdf(PASSPORT), "application/pdf")),
            ("files", ("docs.zip", archive.getvalue(), "application/zip")),
        ],
    )
    bid = client.post(f"/api/v1/packages/{bid['id']}/process").json()
    bid = client.get(f"/api/v1/projects/{pid}").json()["packages"][1]
    assert sorted(d["filename"] for d in bid["documents"]) == [
        "docs.zip/experience/contracts.csv", "passport.pdf", "proposal.pdf"
    ]  # fmt: skip

    run = client.post(f"/api/v1/projects/{pid}/runs", json={"package_id": bid["id"]}).json()
    run = client.get(f"/api/v1/runs/{run['id']}").json()
    assert run["status"] == "done", run["error"]

    matches = {
        m["requirement"]["code"]: m for m in client.get(f"/api/v1/runs/{run['id']}/matches").json()
    }
    assert matches["T-001"]["status"] == "match"  # 16 GB, confirmed by the passport
    assert matches["T-002"]["status"] == "mismatch"  # the passport (official) says 24 months
    assert matches["T-002"]["reasoning"] == "24 months < 36 months"
    assert matches["T-002"]["method"] == "rule"
    assert matches["T-003"]["status"] == "match"
    assert matches["U-001"]["status"] == "mismatch"
    assert matches["U-002"]["status"] == "partial"

    score = run["score"]
    assert score["total"] == 50.0
    assert score["mandatory_done"] == 2 and score["mandatory_total"] == 4
    assert score["verdict"] == "rejection_risk"
    assert score["forecast"] == 100.0
    assert "Участник: Tech LLC" in score["warnings"]

    recs = client.get(f"/api/v1/runs/{run['id']}/recommendations").json()
    warranty = next(r for r in recs if r["code"] == "T-002")
    assert warranty["action_type"] == "replace_product"  # no "fix the number" advice
    assert warranty["impact"] == 21.4
    assert recs[0]["severity"] == "high"

    override = client.patch(
        f"/api/v1/matches/{matches['U-001']['id']}",
        json={"status": "match", "comment": "Сертификат приложен отдельным письмом"},
    ).json()
    assert override["method"] == "manual"
    assert client.get(f"/api/v1/runs/{run['id']}").json()["score"]["total"] == 71.4

    pdf = client.get(f"/api/v1/runs/{run['id']}/report.pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    xlsx = client.get(f"/api/v1/runs/{run['id']}/report.xlsx")
    book = openpyxl.load_workbook(io.BytesIO(xlsx.content))
    assert len(book.sheetnames) == 4
    assert book.worksheets[1]["E3"].value == "Не соответствует"


def test_checklist_is_locked_after_approval(client: TestClient) -> None:
    pid = client.post("/api/v1/projects", json={"title": "Locked"}).json()["id"]
    body = {"scope": "product", "category": "technical", "parameter": "Цвет", "operator": "text",
            "value": "чёрный", "req_type": "desired"}  # fmt: skip
    created = client.post(f"/api/v1/projects/{pid}/requirements", json=body).json()
    assert created["code"] == "T-001" and created["origin"] == "user"

    client.post(f"/api/v1/projects/{pid}/checklist/approve")
    assert client.delete(f"/api/v1/requirements/{created['id']}").status_code == 409

    client.post(f"/api/v1/projects/{pid}/checklist/reopen")
    assert client.delete(f"/api/v1/requirements/{created['id']}").status_code == 204


def test_pipeline_in_one_session(client: TestClient) -> None:
    """Extraction and check share one session in scripts: collections must stay in sync."""
    from app.core.storage import put_bytes
    from app.models import Document, Package, Project, Run
    from app.pipeline.check import run_check
    from app.pipeline.extract import renumber, run_extraction
    from app.pipeline.ingest import normalize_documents

    def add(db: Any, package: Package, name: str, data: bytes) -> None:
        put_bytes(f"t/{package.id}/{name}", data)
        package.documents.append(Document(filename=name, storage_key=f"t/{package.id}/{name}"))
        db.commit()

    with dbmod.SessionLocal() as db:
        project = Project(title="One session")
        db.add(project)
        db.commit()
        customer = Package(project_id=project.id, kind="customer")
        db.add(customer)
        db.commit()
        add(db, customer, "tender.pdf", make_pdf(TENDER))
        normalize_documents(db, customer)
        run_extraction(db, project)
        renumber(project)
        bid = Package(project_id=project.id, kind="bid")
        db.add(bid)
        db.commit()
        add(db, bid, "passport.pdf", make_pdf(PASSPORT))
        normalize_documents(db, bid)
        run = Run(project_id=project.id, package_id=bid.id)
        db.add(run)
        db.commit()
        run_check(db, run)

        assert len(project.requirements) == 5
        assert len(run.matches) == 5
        assert run.score["verdict"] == "rejection_risk"
